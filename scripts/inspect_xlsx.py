# -*- coding: utf-8 -*-
"""查看 xlsx 字频表结构"""
import openpyxl

path = r"C:\dev\self\fuzzy-pinyin-fixer\Chinese character list from 2.5 billion words corpus ordered by frequency.xlsx"
wb = openpyxl.load_workbook(path, read_only=True)
print("sheets:", wb.sheetnames)
ws = wb[wb.sheetnames[0]]
print("dims:", ws.max_row, "rows x", ws.max_column, "cols")
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i < 15:
        print(row)
    else:
        break
